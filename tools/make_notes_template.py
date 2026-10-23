"""Create notes/PLANNING_NOTES.docx, the document Karim fills in before and during planning.

  python tools/make_notes_template.py [--force]

Claude reads it with tools/read_notes.py at the start of the session. Kept out of git (notes/ is ignored).
"""
import pathlib
import sys

from docx import Document
from docx.shared import Pt, RGBColor

REPO = pathlib.Path(__file__).resolve().parents[1]
OUT = REPO / "notes/PLANNING_NOTES.docx"

SECTIONS = [
    ("1. The brief (paste verbatim)", ["Paste the full brief here, exactly as received."]),
    ("2. Rules and logistics", [
        "Time received:            Hard stop (2 h later):",
        "Tools allowed (AI assistants, which LLM APIs, who supplies keys):",
        "Submission format and where to send it:",
        "Public GitHub repo allowed? (yes / no / not stated):",
        "Data provided:",
        "Stated assessment criteria:",
    ]),
    ("3. My first read", [
        "What is being measured (construct):",
        "Who uses the score, for what decision:",
        "What makes this hard / what they are really testing:",
    ]),
    ("4. What I think we should build", ["Thin version that answers the brief:", "What the demo must show in 2 minutes:"]),
    ("5. What we should look at (evidence I want)", ["Psychometric evidence:", "Robustness / fairness checks:"]),
    ("6. Risks: fairness, integrity, candidate experience, regulation", [""]),
    ("7. Assumptions and open questions", [""]),
    ("8. Instructions for Claude (do / avoid)", [""]),
]


def main():
    if OUT.exists() and "--force" not in sys.argv:
        sys.exit(f"{OUT.relative_to(REPO)} already exists (use --force to overwrite)")
    OUT.parent.mkdir(exist_ok=True)
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Rubik"; st.font.size = Pt(10.5)
    t = doc.add_heading("Work sample: planning notes", level=0)
    for r in t.runs:
        r.font.name = "Rubik Medium"; r.font.color.rgb = RGBColor.from_string("22252B")
    doc.add_paragraph("Private working notes (not submitted, not in git). Claude reads this first when planning.")
    for head, prompts in SECTIONS:
        doc.add_heading(head, level=1)
        for p in prompts:
            doc.add_paragraph(p)
    doc.save(OUT)
    print(f"Wrote {OUT.relative_to(REPO)}")


if __name__ == "__main__":
    main()
